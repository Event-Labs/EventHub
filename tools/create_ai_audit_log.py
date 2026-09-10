"""Create an EventHub AI Audit Log workbook from the supplied Excel template.

Run:
    python tools/create_ai_audit_log.py

Optional:
    python tools/create_ai_audit_log.py --student-id DE123456
"""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


def reflection(critical: str, context: str, synthesis: str, decision: str) -> str:
    return (
        f"Critical Thinking: {critical}\n"
        f"Contextualization: {context}\n"
        f"Creative Synthesis: {synthesis}\n"
        f"Decision: {decision}"
    )


ENTRIES = [
    (
        "DECISION", "Decomposition",
        "Tổ chức toàn bộ khu vực Staff theo các nghiệp vụ vận hành sự kiện.",
        "Hãy xây dựng khu vực Staff thành các màn hình riêng cho dashboard, sự kiện, check-in, báo cáo và đặt vé.",
        "AI đề xuất gộp dashboard, sự kiện, check-in và bán vé vào một trang.",
        reflection(
            "Một trang lớn khó sử dụng trên thiết bị di động và khó quản lý trạng thái.",
            "Staff cần dashboard, sự kiện, QR, thủ công, báo cáo, nhiệm vụ và đặt vé.",
            "Dùng StaffLayout chung nhưng tách route/page theo từng nghiệp vụ.",
            "Giữ navigation tập trung và mỗi chức năng có component, service riêng.",
        ),
        "frontend/src/pages/staff/StaffLayout.jsx; frontend/src/pages/staff/StaffDashboardPage.jsx",
    ),
    (
        "DECISION", "Abstraction",
        "Luồng QR cần tách bước đọc/kiểm tra vé và bước xác nhận check-in.",
        "Hãy tách bước quét và xác thực QR khỏi bước staff xác nhận check-in.",
        "AI đề xuất camera đọc được mã là cập nhật check-in ngay.",
        reflection(
            "Tự động cập nhật dễ check-in nhầm khi camera đọc lặp hoặc staff chưa xem thông tin.",
            "Giao diện có bước verify, hiển thị người mua, sự kiện và trạng thái vé.",
            "Tách POST verify-qr khỏi POST check-in/qr để staff xác nhận.",
            "Chỉ gọi check-in sau khi verify thành công và người dùng bấm xác nhận.",
        ),
        "backend/src/modules/tickets/staffTickets.routes.js; frontend/src/pages/staff/StaffQrCheckInPage.jsx",
    ),
    (
        "VERIFICATION", "Algorithms",
        "Xác thực nội dung QR để tránh dùng ticketId bị sửa hoặc QR của vé khác.",
        "Hãy xác thực toàn bộ QR payload, đối chiếu với vé và kiểm tra quyền của staff.",
        "AI nói ticketRef tồn tại là đủ xác nhận QR hợp lệ.",
        reflection(
            "ẢO GIÁC: ticketRef đúng chưa chứng minh toàn bộ QR khớp với vé.",
            "Payload QR có thể bị chỉnh sửa và staff chỉ được truy cập event được phân công.",
            "Parse payload, kiểm tra quyền, thời gian và dùng assertQrMatchesTicket.",
            "Từ chối QR sai định dạng, sai vé hoặc ngoài phạm vi staff.",
        ),
        "backend/src/modules/tickets/tickets.service.js: staffVerifyTicketByQr/assertQrMatchesTicket",
    ),
    (
        "PROBLEM-SOLVING", "Pattern Recognition",
        "Không cho check-in trước thời gian do organizer cấu hình.",
        "Hãy chặn check-in trước giờ mở cổng ở backend và hiển thị trạng thái tương ứng trên frontend.",
        "AI đề xuất chỉ disable nút quét trước checkin_start_time.",
        reflection(
            "Frontend có thể bị bỏ qua bằng cách gọi API trực tiếp.",
            "StaffEventsPage cần thông báo thời gian, nhưng backend mới là nguồn kiểm soát.",
            "Dùng assertCheckInOpen ở service và hiển thị countdown/trạng thái ở UI.",
            "Chặn tại backend, đồng thời hướng dẫn staff chờ đúng giờ trên frontend.",
        ),
        "backend/src/modules/tickets/tickets.service.js: assertCheckInOpen; frontend/src/pages/staff/StaffEventsPage.jsx",
    ),
    (
        "PROBLEM-SOLVING", "Algorithms",
        "Chống hai thiết bị check-in cùng một vé tại nhiều cổng.",
        "Hãy triển khai cập nhật check-in nguyên tử để chống hai thiết bị quét trùng cùng một vé.",
        "AI đề xuất SELECT trạng thái rồi UPDATE ở câu lệnh tiếp theo.",
        reflection(
            "ẢO GIÁC: hai request có thể cùng đọc trạng thái chưa dùng trước khi update.",
            "Sự kiện có thể có nhiều staff quét QR đồng thời.",
            "Dùng conditional update/transaction và trả state cụ thể từ repository.",
            "Chỉ một request thành công; request sau nhận trạng thái vé đã check-in.",
        ),
        "backend/src/modules/tickets/tickets.repository.js: checkInTicket; tickets.service.js: staffCheckInTicket",
    ),
    (
        "DECISION", "Decomposition",
        "Check-in thủ công cần chọn sự kiện trước rồi tìm vé.",
        "Hãy xây dựng luồng check-in thủ công theo mã vé, email hoặc số điện thoại sau khi chọn sự kiện.",
        "AI đề xuất ô tìm kiếm toàn hệ thống không cần chọn event.",
        reflection(
            "Tìm toàn hệ thống dễ lộ dữ liệu và chọn nhầm vé sự kiện khác.",
            "Service yêu cầu eventId và repository giới hạn theo assignment của staff.",
            "Chia luồng chọn event, nhập từ khóa, xem danh sách và chọn vé.",
            "Bắt buộc eventId trước khi gọi staffSearchTickets.",
        ),
        "backend/src/modules/tickets/tickets.service.js: staffSearchTickets; frontend/src/pages/staff/StaffQrCheckInPage.jsx",
    ),
    (
        "DECISION", "Abstraction",
        "Staff cần xem đúng thông tin trước khi xác nhận check-in thủ công.",
        "Hãy hiển thị đầy đủ thông tin vé và khách hàng trước khi staff xác nhận check-in thủ công.",
        "AI đề xuất chỉ hiển thị ticket_code và nút Check-in.",
        reflection(
            "Chỉ có mã vé không đủ để staff đối chiếu khách và sự kiện.",
            "Cần tên người mua, loại vé, event, trạng thái và thời điểm check-in.",
            "Chuẩn hóa buildStaffTicketPayload dùng chung cho QR và manual.",
            "Hiển thị detail panel và yêu cầu xác nhận trước PATCH check-in.",
        ),
        "backend/src/modules/tickets/tickets.service.js: buildStaffTicketPayload/getStaffTicket; frontend/src/pages/staff/StaffQrCheckInPage.jsx",
    ),
    (
        "VERIFICATION", "Pattern Recognition",
        "Phân loại lỗi check-in để staff biết cách xử lý tại cổng.",
        "Hãy phân loại và hiển thị riêng từng lỗi check-in để staff biết cách xử lý.",
        "AI đề xuất dùng một thông báo lỗi chung cho mọi trường hợp.",
        reflection(
            "Thông báo chung khiến staff không biết vé chưa trả tiền, hết hạn hay sai quyền.",
            "Repository trả NOT_FOUND, FORBIDDEN, INVALID_STATUS, INVALID_ORDER, CHECKIN_NOT_OPEN và EXPIRED.",
            "Map từng state thành mã HTTP và thông báo riêng trong service.",
            "Hiển thị modal kết quả phù hợp để staff xử lý nhanh.",
        ),
        "backend/src/modules/tickets/tickets.service.js: staffCheckInTicket; frontend/src/pages/staff/StaffResultsPage.jsx",
    ),
    (
        "PROBLEM-SOLVING", "Algorithms",
        "Trang thống kê cần tổng vé, đã vào, còn lại, theo giờ và lượt gần nhất.",
        "Hãy tạo báo cáo check-in gồm tổng vé, đã vào, còn lại, phân bố theo giờ và lượt gần nhất.",
        "AI đề xuất tải tất cả vé về browser rồi group bằng JavaScript.",
        reflection(
            "Dữ liệu lớn làm chậm trình duyệt và tăng lưu lượng mạng.",
            "UI chỉ cần số tổng hợp và danh sách check-in gần đây.",
            "Dùng aggregate COUNT/GROUP BY tại repository và giới hạn recent rows.",
            "Trả một check-in report cho trang thống kê và chi tiết event.",
        ),
        "backend/src/modules/operations/operations.repository.js; frontend/src/pages/staff/StaffCheckInCountPage.jsx",
    ),
    (
        "DECISION", "Decomposition",
        "Chỉ hiển thị sự kiện mà staff được phép bán vé trực tiếp.",
        "Hãy giới hạn danh sách sự kiện đặt vé trực tiếp theo assignment và role của staff.",
        "AI đề xuất trả mọi sự kiện đang published cho bất kỳ STAFF nào.",
        reflection(
            "Role STAFF toàn cục không đồng nghĩa được vận hành mọi sự kiện.",
            "Quyền phải dựa trên assignment; ADMIN/SUPER_ADMIN có phạm vi rộng hơn.",
            "Truyền staffId và roles vào findStaffDirectBookingEvents.",
            "Chỉ trả event và ticket types mà tài khoản được phép thao tác.",
        ),
        "backend/src/modules/orders/orders.service.js: getStaffDirectBookingEvents; orders.routes.js",
    ),
    (
        "PROBLEM-SOLVING", "Algorithms",
        "Đặt vé trực tiếp phải giữ ghế và tạo order/items nhất quán.",
        "Hãy dùng một transaction checkout chung cho đặt vé trực tiếp có ghế và không có ghế.",
        "AI đề xuất tạo hai luồng SQL riêng cho có ghế và không có ghế.",
        reflection(
            "Hai luồng dễ lệch schema, tính tiền và kiểm soát tồn kho.",
            "Dự án hỗ trợ ticket hold, session seats và nhiều phương thức thanh toán.",
            "Dùng chung createPendingCheckout transaction cho CASH và bank_transfer.",
            "Một checkout path xử lý item, ghế, hold và order nhất quán.",
        ),
        "backend/src/modules/orders/orders.service.js: createStaffDirectBooking/createPendingCheckout",
    ),
    (
        "DECISION", "Pattern Recognition",
        "Thanh toán tiền mặt tại quầy cần xác nhận và phát hành vé ngay.",
        "Hãy hoàn tất thanh toán CASH, phát hành vé và gửi email ngay sau khi staff xác nhận nhận tiền.",
        "AI đề xuất tạo order PENDING rồi chờ webhook giống PayOS.",
        reflection(
            "Tiền mặt không có webhook, để PENDING sẽ khiến staff không thể in vé.",
            "Staff là người xác nhận đã nhận tiền tại quầy.",
            "Confirm manual payment, phát hành vé, gửi email rồi đọc trạng thái cuối.",
            "CASH hoàn tất ngay và trả confirmation_email_sent cho giao diện.",
        ),
        "backend/src/modules/orders/orders.service.js: confirmStaffDirectManualPayment/sendTicketConfirmation",
    ),
    (
        "VERIFICATION", "Decomposition",
        "PayOS tại quầy phải giữ order PENDING cho đến khi thanh toán thật.",
        "Hãy giữ đơn PayOS ở trạng thái PENDING cho đến khi webhook hoặc status sync xác nhận PAID.",
        "AI nói tạo được checkout_url đồng nghĩa khách đã thanh toán.",
        reflection(
            "ẢO GIÁC: checkout_url chỉ là yêu cầu thanh toán, không phải bằng chứng đã trả tiền.",
            "Backend có payment channel, PayOS link, return/cancel URL và đồng bộ trạng thái.",
            "Tạo pending checkout, tạo link; nếu tạo link lỗi thì cancel order.",
            "Chỉ phát hành/in vé sau khi webhook hoặc sync xác nhận PAID.",
        ),
        "backend/src/modules/orders/orders.service.js: createTicketOrderPayosLink/getStaffDirectBookingStatus",
    ),
    (
        "PROBLEM-SOLVING", "Abstraction",
        "Frontend cần theo dõi PayOS nhưng không cấp vé dựa trên thời gian chờ.",
        "Hãy polling trạng thái đơn PayOS và chỉ bật chức năng in vé khi đơn đã PAID.",
        "AI đề xuất chờ cố định 10 giây rồi chuyển giao diện sang thành công.",
        reflection(
            "Thời gian cố định có thể báo sai khi khách chưa trả hoặc PayOS chậm.",
            "Service có endpoint status và cơ chế sync PayOS theo orderId.",
            "Polling khi PENDING, dừng khi PAID/CANCELLED/EXPIRED và có nút refresh.",
            "Chỉ bật in vé khi displayResult.order.status là PAID.",
        ),
        "frontend/src/pages/staff/StaffDirectBookingPage.jsx: statusQuery/BookingResult; frontend/src/services/orders.js",
    ),
    (
        "DECISION", "Pattern Recognition",
        "Sau đặt vé, staff cần in vé và khách vẫn nhận vé qua email.",
        "Hãy xây dựng màn hình kết quả đặt vé trực tiếp có thông tin đơn, thanh toán, vé, email và chức năng in.",
        "AI đề xuất luôn hiển thị nút in ngay khi order được tạo.",
        reflection(
            "In khi PENDING có thể phát hành vé chưa thanh toán.",
            "Kết quả gồm order, buyer, items, payment, tickets và trạng thái gửi email.",
            "Tạo canPrint = PAID và tickets.length > 0; render QR riêng cho từng vé.",
            "Chỉ cho in sau PAID, đồng thời thông báo kết quả gửi email.",
        ),
        "frontend/src/pages/staff/StaffDirectBookingPage.jsx: BookingResult/direct-booking-print",
    ),
    (
        "PROBLEM-SOLVING", "Abstraction",
        "Trình duyệt có thể từ chối quyền camera khi staff quét QR.",
        "Hãy xử lý quyền camera bị từ chối và cung cấp cách chuyển sang check-in thủ công.",
        "AI đề xuất tải lại trang liên tục cho đến khi trình duyệt cấp quyền.",
        reflection(
            "Tải lại không giải quyết quyền bị chặn và làm staff mất thao tác đang thực hiện.",
            "Trang QR đã có trạng thái camera denied và luồng manual check-in thay thế.",
            "Hiển thị hướng dẫn cấp quyền, nút thử lại và liên kết sang tìm vé thủ công.",
            "Dùng CameraDeniedContent thay vì vòng lặp reload.",
        ),
        "frontend/src/pages/staff/StaffQrCheckInPage.jsx: CameraDeniedContent/CameraErrorContent",
    ),
    (
        "PROBLEM-SOLVING", "Algorithms",
        "Camera có thể đọc cùng một QR nhiều lần trong vài mili giây.",
        "Hãy chống callback quét QR lặp và dừng camera trong lúc xác thực vé.",
        "AI đề xuất gửi API cho mọi kết quả mà camera đọc được.",
        reflection(
            "Nhiều callback gây request trùng, modal nhấp nháy và tăng nguy cơ double check-in.",
            "QR scanner chạy liên tục trong khi API verify cần thời gian phản hồi.",
            "Khóa trạng thái processing theo QR gần nhất và pause scanner trước request.",
            "Chỉ resume camera sau khi staff đóng kết quả hoặc quét vé khác.",
        ),
        "frontend/src/pages/staff/StaffQrCheckInPage.jsx: scanner callbacks/checkInState",
    ),
    (
        "DECISION", "Pattern Recognition",
        "Staff cần cập nhật tiến độ các nhiệm vụ được organizer giao.",
        "Hãy xây dựng trang nhiệm vụ Staff có lọc trạng thái và cập nhật tiến độ an toàn.",
        "AI đề xuất cho phép staff chỉnh sửa mọi nhiệm vụ nhìn thấy trên giao diện.",
        reflection(
            "Hiển thị được task không đồng nghĩa có quyền cập nhật task của người khác.",
            "API nhận staff hiện tại và taskId, trạng thái cần nằm trong tập hợp hợp lệ.",
            "Kiểm tra ownership ở service và cập nhật UI sau khi server xác nhận.",
            "Dùng fetchAssignedStaffTasks và updateAssignedStaffTaskStatus.",
        ),
        "frontend/src/pages/staff/StaffTasksPage.jsx; frontend/src/services/operations.js",
    ),
    (
        "DECISION", "Decomposition",
        "Form đặt vé trực tiếp cần thu thập thông tin người mua và nhiều loại vé.",
        "Hãy chia form đặt vé trực tiếp thành bước chọn sự kiện, chọn vé, nhập khách hàng và chọn thanh toán.",
        "AI đề xuất đặt toàn bộ trường vào một form dài và gửi thẳng lên API.",
        reflection(
            "Form dài khó kiểm tra và dễ gửi item rỗng hoặc thông tin khách thiếu.",
            "Luồng có event, ticket types, session seats, buyer và payment method.",
            "Tách step, validate từng phần và tạo payload chuẩn trước khi submit.",
            "Chỉ cho tạo đơn khi có ít nhất một item hợp lệ và buyer hợp lệ.",
        ),
        "frontend/src/pages/staff/StaffDirectBookingPage.jsx; backend/src/modules/orders/orders.validation.js",
    ),
    (
        "PROBLEM-SOLVING", "Algorithms",
        "Số lượng vé và danh sách ghế được chọn phải đồng nhất.",
        "Hãy kiểm tra quantity, session_seat_ids và tồn kho trước khi tạo đơn trực tiếp.",
        "AI đề xuất tin số lượng và ghế do frontend gửi lên.",
        reflection(
            "Payload có thể bị sửa và hai staff có thể cùng chọn một ghế.",
            "Backend có ticket holds, session seats và transactional checkout.",
            "Chuẩn hóa quantity, kiểm tra số ghế, khóa/giữ ghế và xác nhận tồn kho trong transaction.",
            "Từ chối item không hợp lệ hoặc ghế đã được giữ/bán.",
        ),
        "backend/src/modules/orders/orders.service.js: normalizedItems/createPendingCheckout",
    ),
    (
        "PROBLEM-SOLVING", "Pattern Recognition",
        "Đơn PayOS có thể bị hủy, hết hạn hoặc lỗi đồng bộ trong lúc staff chờ.",
        "Hãy xử lý đầy đủ trạng thái PENDING, PAID, CANCELLED và EXPIRED trên màn hình đặt vé trực tiếp.",
        "AI đề xuất chỉ xử lý PENDING và PAID.",
        reflection(
            "Bỏ qua trạng thái cuối khác làm giao diện polling mãi và giữ ghế không cần thiết.",
            "Backend gọi expirePendingOrders và status sync có thể tạm thời thất bại.",
            "Dừng polling ở mọi terminal state, hiển thị hướng dẫn và cho tạo đơn mới khi cần.",
            "Giữ status endpoint sử dụng được ngay cả khi PayOS sync tạm lỗi.",
        ),
        "backend/src/modules/orders/orders.service.js: getStaffDirectBookingStatus; frontend/src/pages/staff/StaffDirectBookingPage.jsx",
    ),
]


HALLUCINATIONS = [
    (
        "003", "Security Oversimplification",
        "Có ticketRef tồn tại là đủ kết luận QR hợp lệ.",
        "Phải kiểm tra payload khớp vé, quyền staff và thời gian check-in.",
        "Đối chiếu staffVerifyTicketByQr và thử payload đã sửa.",
        "Dùng extractQrTicketPayload, assertQrMatchesTicket và access check.",
    ),
    (
        "005", "Concurrency Error",
        "SELECT rồi UPDATE riêng biệt vẫn chống được hai cổng quét trùng.",
        "Hai request có thể cùng đọc trạng thái chưa check-in trước khi update.",
        "Mô phỏng hai request đồng thời cho cùng ticketId.",
        "Dùng conditional update/transaction và trả state đã check-in.",
    ),
    (
        "013", "Context Misunderstanding",
        "Tạo được PayOS checkout URL nghĩa là đơn đã thanh toán.",
        "Checkout URL chỉ mở giao dịch; order phải giữ PENDING đến khi PayOS xác nhận.",
        "Đối chiếu payment status trước và sau webhook/status sync.",
        "Chỉ phát hành vé khi trạng thái cuối là PAID.",
    ),
]


def set_values(ws, values: dict[str, object]) -> None:
    for cell, value in values.items():
        ws[cell] = value


def copy_row_style(ws, source_row: int, target_row: int, max_column: int) -> None:
    """Copy the template formatting when the audit log grows past its original rows."""
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def create_workbook(
    template: Path,
    output: Path,
    student_name: str,
    student_id: str,
    course: str,
) -> None:
    workbook = load_workbook(template)

    summary = workbook["1. Metadata & Summary"]
    set_values(summary, {
        "C4": student_name,
        "C5": student_id,
        "C6": course,
        "C7": "EventHub - Event Management Platform",
        "C10": 105,
        "C11": 21,
        "C12": 0.2,
        "C13": 3,
        "A17": "ChatGPT / Codex",
        "B17": "Thiết kế, sinh mã, debugging và review",
        "C17": "High",
        "D17": "Tăng tốc triển khai và phản biện",
        "A18": "GitHub Copilot",
        "B18": "Gợi ý code và hoàn thiện cú pháp",
        "C18": "Medium",
        "D18": "Giảm thao tác lặp",
        "A19": "Manual verification",
        "B19": "Đọc code, kiểm thử và đối chiếu nghiệp vụ",
        "C19": "High",
        "D19": "Xác nhận tính đúng đắn",
        "A20": None, "B20": None, "C20": None, "D20": None,
        "B24": 5, "B25": 6, "B26": 4, "B27": 6,
    })

    detail = workbook["2. Detailed Audit Log"]
    for index, entry in enumerate(ENTRIES, start=1):
        row = index + 3
        values = (f"{index:03d}",) + entry
        for column, value in enumerate(values, start=1):
            detail.cell(row=row, column=column, value=value)

    hallucination = workbook["3. Hallucination Detection"]
    for row_index, values in enumerate(HALLUCINATIONS, start=4):
        for column, value in enumerate(values, start=1):
            hallucination.cell(row=row_index, column=column, value=value)

    checklist = workbook["4. Self-Assessment Checklist"]
    set_values(checklist, {
        "D14": "21 entries (theo yêu cầu cập nhật)",
        "D15": "Đầy đủ 4 DTC components",
        "D16": "3 cases (đạt yêu cầu Project)",
        "D17": "Tất cả 21 entries đầy đủ 4 phần Human Delta",
        "D18": "100% entries có evidence trong mã nguồn",
    })
    for index, entry in enumerate(ENTRIES, start=1):
        row = index + 27
        if row > 42:
            copy_row_style(checklist, 42, row, 4)
        decision = entry[5].splitlines()[-1].replace("Decision: ", "")
        checklist.cell(row, 1, f"{index:03d}")
        checklist.cell(row, 2, decision)
        checklist.cell(row, 3, entry[4])
        checklist.cell(row, 4, entry[6])

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()

    # A second load verifies that the generated file is a readable XLSX workbook.
    verified = load_workbook(output, read_only=True, data_only=False)
    expected_sheets = {
        "1. Metadata & Summary",
        "2. Detailed Audit Log",
        "3. Hallucination Detection",
        "4. Self-Assessment Checklist",
    }
    if set(verified.sheetnames) != expected_sheets:
        raise RuntimeError("Generated workbook does not contain the expected sheets")
    detail_check = verified["2. Detailed Audit Log"]
    if detail_check["A24"].value != "021":
        raise RuntimeError("Generated workbook is missing audit-log entries")
    prompt_values = [detail_check.cell(row, 5).value or "" for row in range(4, 25)]
    if any("?" in prompt for prompt in prompt_values):
        raise RuntimeError("Prompt to AI must contain instructions, not questions")
    verified.close()


def parse_args() -> argparse.Namespace:
    repo = Path(__file__).resolve().parents[1]
    audit_dir = repo.parents[1]
    parser = argparse.ArgumentParser(description="Create EventHub AI Audit Log")
    parser.add_argument(
        "--template",
        type=Path,
        default=audit_dir / "AI_AuditLog_BuiLeLongDai_DE180691_WDP301_SE18D02.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=audit_dir / "AI_AuditLog_EventHub_Staff_21_Prompts.xlsx",
    )
    parser.add_argument("--student-name", default="Bùi Tuấn Khánh")
    parser.add_argument("--student-id", default="Chưa cung cấp")
    parser.add_argument("--course", default="WDP301")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    create_workbook(
        template=args.template.resolve(),
        output=args.output.resolve(),
        student_name=args.student_name,
        student_id=args.student_id,
        course=args.course,
    )
    print(f"Created and verified: {args.output.resolve()}")
