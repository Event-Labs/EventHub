import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'TestCases'

# Visual styles matching the exact screenshot
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
func_fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid') # light blue banner
header_font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
func_font = Font(name='Arial', size=10, bold=True, color='000000')
data_font = Font(name='Arial', size=9, bold=False, color='000000')

thin_side = Side(border_style='thin', color='7F7F7F')
border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

headers = [
    "Test Case ID",
    "Test Case Description",
    "Test Steps",
    "Expected Result",
    "Pre-condition",
    "Execution 1",
    "Tester",
    "Test Date",
    "Execution 2",
    "Tester",
    "Test Date",
    "Execution 3",
    "Tester",
    "Test Date"
]

test_cases_func6 = [
    (
        "TC-AUM-VPF-001",
        "Test viewing the user's profile information.",
        "1. Login successfully.\n2. Open Profile.\n3. Observe the profile information.",
        "Profile page displays the user's available information, including full name, email, phone number, avatar, and role.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-002",
        "Test viewing profile with different account roles.",
        "1. Login as Customer/Organizer/Staff/Admin.\n2. Open Profile.\n3. Check displayed role information.",
        "System displays the correct account information and role for the logged-in user.",
        "User accounts with different roles exist.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-003",
        "Test accessing profile without authentication.",
        "1. Log out or clear session.\n2. Attempt to open profile page directly.\n3. Observe system response.",
        "System rejects access and displays session expiration or redirects to Login page.",
        "User is not authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-004",
        "Test viewing profile with empty optional fields.",
        "1. Login with account having null phone/address/dob.\n2. Open Profile.\n3. Observe unfilled fields.",
        "System displays placeholder 'Chưa cập nhật' and default letter avatar without UI errors.",
        "User profile has null optional fields.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-005",
        "Test viewing Organizer profile when 2FA is disabled.",
        "1. Login as Organizer with 2FA disabled.\n2. Open Organizer Profile.\n3. Observe gatekeeper screen.",
        "System presents 2FA Gatekeeper requiring user to set up 2FA via email OTP to view organizer profile.",
        "User is authenticated as Organizer with 2FA disabled.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-006",
        "Test viewing Organizer sensitive legal info & verification documents.",
        "1. Open Organizer Profile (2FA active).\n2. Click locked Legal / Documents section.\n3. Enter 6-digit OTP sent to email.",
        "System validates OTP and unlocks sensitive legal info and documents for 10 minutes.",
        "Organizer is authenticated with 2FA enabled.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-007",
        "Test viewing Admin profile, security score, and sessions.",
        "1. Login as Admin.\n2. Open Admin Profile (/admin/profile).\n3. Check credentials, security audit, and login sessions.",
        "Admin profile displays administrator details, security health score, 2FA status, and recent login sessions.",
        "User is authenticated as Admin.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-VPF-008",
        "Test profile loading error handling and retry mechanism.",
        "1. Simulate network disconnect or 500 error.\n2. Open Profile.\n3. Reconnect and click Retry button.",
        "System displays an error alert with a Retry button, successfully reloading data on click.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    )
]

test_cases_func7 = [
    (
        "TC-AUM-UPF-001",
        "Test updating profile with valid information.",
        "1. Login successfully.\n2. Open Profile.\n3. Click Edit/Update Profile.\n4. Modify full name/phone/avatar.\n5. Click Save.",
        "Profile is updated successfully and the new information is displayed.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-002",
        "Test updating profile with invalid phone number.",
        "1. Open Update Profile.\n2. Enter an invalid phone number.\n3. Click Save.",
        "System rejects invalid data and displays an appropriate validation message. Existing valid profile information remains unchanged.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-003",
        "Test cancelling profile changes.",
        "1. Open Update Profile.\n2. Modify profile information.\n3. Click Cancel/Back instead of Save.\n4. Reopen Profile.",
        "Changes are discarded and the original profile information remains unchanged.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-004",
        "Test updating profile with international phone number (+84).",
        "1. Open Update Profile.\n2. Enter phone with +84 format (e.g. +84912345678).\n3. Click Save.",
        "System accepts and normalizes phone number to standard 09xxxxxxxx format and saves successfully.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-005",
        "Test updating profile with empty full name.",
        "1. Open Update Profile.\n2. Clear Full Name field.\n3. Click Save.",
        "System rejects submission and displays validation error 'Vui lòng nhập họ và tên.'",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-006",
        "Test uploading valid image file as avatar.",
        "1. Open Update Profile.\n2. Choose a valid image file (JPG, PNG, WEBP).\n3. Click Save.",
        "Image is previewed immediately, uploaded to Cloudinary, and saved as user's avatar.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-007",
        "Test uploading non-image file as avatar.",
        "1. Open Update Profile.\n2. Choose a non-image file (PDF, TXT, DOCX).\n3. Observe feedback.",
        "System rejects file immediately with message 'Vui lòng chọn tệp ảnh hợp lệ (JPG, PNG).'",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-008",
        "Test verifying Account Email is read-only.",
        "1. Open Update Profile.\n2. Check email input field.",
        "Email field is disabled / fixed and cannot be edited by the user.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-009",
        "Test clearing optional profile fields.",
        "1. Open Update Profile for user with existing optional data.\n2. Clear phone, DOB, city, and address.\n3. Click Save.",
        "Profile is updated successfully, setting optional fields to null without errors.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-010",
        "Test submitting Organizer verification profile update request.",
        "1. Open Organizer Profile -> Legal and Documents.\n2. Click 'Yêu cầu cập nhật'.\n3. Enter legal details & upload verification files.\n4. Click Submit.",
        "System submits update request with status PENDING for Admin review and displays success toast.",
        "Organizer is authenticated with sensitive access unlocked.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-011",
        "Test duplicate Organizer update request prevention.",
        "1. Login as Organizer with a PENDING update request.\n2. Open Legal & Documents section.\n3. Check update button.",
        "Update button is disabled with banner stating a request is already pending Admin review.",
        "Organizer has a pending profile update request.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-UPF-012",
        "Test handling network error during profile save.",
        "1. Open Update Profile and modify information.\n2. Disconnect network and click Save.\n3. Observe system behavior.",
        "System displays error message 'Không thể cập nhật hồ sơ.', keeping entered data intact.",
        "User is authenticated.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    )
]

test_cases_func12 = [
    (
        "TC-AUM-REA-001",
        "Test viewing empty AI review state before analysis.",
        "1. Login as Admin.\n2. Open Admin Event Review Detail for an unanalyzed PENDING event.\n3. Check AI Assistant right panel.",
        "System displays empty state banner 'Chưa chạy AI Đánh giá' with description and button 'Bắt đầu Phân tích AI ngay'.",
        "User is authenticated as Admin; event is pending review without prior AI review.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-002",
        "Test running AI Assistant analysis on pending event.",
        "1. Click 'Bắt đầu Phân tích AI ngay' or 'Chạy lại AI'.\n2. Observe active scanning progress state.\n3. Wait for analysis completion.",
        "System displays multi-phase scanning pulse state (Vision image check & Qwen3 policy match), completes analysis, and renders full AI report with toast 'Đã hoàn thành phân tích sự kiện bằng AI!'.",
        "User is authenticated as Admin; event is in PENDING_REVIEW status.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-003",
        "Test AI recommendation for compliant event (APPROVE).",
        "1. Open review detail for an event meeting all policies.\n2. Run or observe AI Assistant assessment.",
        "AI displays green banner 'Đủ điều kiện phê duyệt' (APPROVE), low risk score (<=20), high quality score (>=80), and lists all passed compliant items.",
        "Event data is fully compliant with platform policies.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-004",
        "Test AI recommendation for high-risk violation event (REJECT).",
        "1. Open review detail for an event containing prohibited or sensitive content.\n2. Observe AI assessment.",
        "AI displays red banner 'Nguy cơ vi phạm nghiêm trọng' (REJECT), high risk score (>50), and highlights critical policy violations with quoted text.",
        "Event contains policy-violating content.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-005",
        "Test AI recommendation for borderline event (NEEDS_REVIEW).",
        "1. Open review detail for an event with missing optional licenses or unclear clauses.\n2. Observe AI assessment.",
        "AI displays amber banner 'Cần quản trị viên đối chiếu' (NEEDS_REVIEW), medium risk score (20-50), and lists yellow caution warnings.",
        "Event has ambiguous terms requiring human discretion.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-006",
        "Test AI Quality Score and Risk Index meters.",
        "1. Observe 'Chất lượng thông tin' (Quality Score) and 'Chỉ số rủi ro' (Risk Index) meters.\n2. Check progress bar colors and numeric values.",
        "System renders visual progress bars: Quality Score (green gradient) and Risk Score (green/amber/rose gradient based on risk severity).",
        "AI review has been executed.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-007",
        "Test AI Title and Description quality audit.",
        "1. Check 'Soát lỗi Tiêu đề, Mô tả & Chính tả' card in AI panel.\n2. Observe title and description status badges and analysis notes.",
        "System displays status badges (Hợp lệ / Cần sửa / Vi phạm) and detailed analysis notes for both Event Title and Event Description.",
        "AI review has completed.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-008",
        "Test AI Vietnamese spelling and grammar detection.",
        "1. Open review detail for an event containing spelling and grammatical errors.\n2. Check spelling issues section.",
        "AI detects each spelling/grammar mistake with bullet points and displays '+ Thêm' button next to each error.",
        "Event description contains spelling mistakes.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-009",
        "Test AI Vision & OCR Image Moderation on Poster and Banner.",
        "1. Check 'Kiểm duyệt Hình ảnh & OCR' section in AI panel.\n2. Observe per-image status badges and analysis.",
        "AI classifies each image (Main Poster, Cover Banner, Description Images) with verdicts (Hợp lệ / Cần lưu ý / Vi phạm / Chưa có) and displays OCR text/safety analysis.",
        "Event has uploaded image assets.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-010",
        "Test single-click insertion of individual AI findings into review notes.",
        "1. Click on a Policy Violation, Warning, Spelling Issue, or Suggestion card.\n2. Check the 'Ghi chú & Quyết định kiểm duyệt' textarea.",
        "Selected AI finding is automatically appended with prefix tag (e.g. `[Vi phạm ...]`, `[Lỗi chính tả] ...`) into review notes textarea with toast confirmation.",
        "AI findings are visible on screen.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-011",
        "Test 'Sao chép tất cả lỗi' (Copy All Issues) into review notes.",
        "1. Click 'Sao chép tất cả lỗi' button at the top of AI analysis details.\n2. Observe review notes textarea.",
        "System aggregates all critical violations, warnings, and adjustment suggestions into structured sections and inserts them into the review notes textarea.",
        "AI review contains multiple issues.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-012",
        "Test previewing attached event permit documents (DocumentPreviewModal).",
        "1. Under 'Giấy phép & Tài liệu đính kèm', click 'Xem' on a PDF or image permit.\n2. Check Document Preview Modal.",
        "Modal opens displaying document preview, filename, file type, file size, 'Mở tab mới' link, and 'Tải về máy' download button.",
        "Event has attached legal permit files.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-013",
        "Test closing Document Preview Modal via Close button and ESC key.",
        "1. Open Document Preview Modal.\n2. Press Escape key on keyboard (or click backdrop / close button).\n3. Observe modal behavior.",
        "Modal closes smoothly and returns focus to the Event Review Detail page.",
        "Document Preview Modal is currently open.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-014",
        "Test inspecting event sessions, ticket tiers, and refund policy details.",
        "1. Scroll through event details column.\n2. Verify Sessions schedule, Ticket types (name, quantity, price), and Refund Policy terms.",
        "All event parameters, session timings, ticket pricing, and attached refund policy documents are accurately rendered.",
        "Event details are loaded.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-015",
        "Test approving event with confirmation modal.",
        "1. Click 'Phê duyệt' button.\n2. In Confirm Modal, verify prompt and click 'Phê duyệt'.\n3. Observe redirection and status.",
        "Confirm modal opens; upon confirmation, event is approved, toast 'Đã phê duyệt sự kiện.' is shown, and admin is navigated to review list.",
        "Event is in PENDING_REVIEW status.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-016",
        "Test rejecting event with review notes justification.",
        "1. Enter or insert rejection reasons into review notes.\n2. Click 'Từ chối' button.\n3. Confirm rejection in Confirm Modal.",
        "Confirm modal opens with red danger button; upon confirmation, event is rejected, status updates to REJECTED, and reasons are logged.",
        "Event is in PENDING_REVIEW status.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-017",
        "Test viewing non-pending event detail (APPROVED / REJECTED / HIDDEN).",
        "1. Open an event with status APPROVED or REJECTED.\n2. Observe page layout.",
        "Right AI review/action sidebar is hidden; left column expands to full width displaying event details and existing admin review note banner.",
        "Event status is not PENDING_REVIEW.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    ),
    (
        "TC-AUM-REA-018",
        "Test error handling during AI service failure or timeout.",
        "1. Disconnect AI service or simulate API 500/timeout.\n2. Click 'Chạy lại AI'.\n3. Observe error handling and manual review availability.",
        "System displays toast 'Không thể chạy phân tích AI.', page remains fully responsive, and manual Approve/Reject buttons remain accessible.",
        "AI service encounters an error.",
        "Pending", "", "", "Pending", "", "", "Pending", "", ""
    )
]

total_cols = len(headers)

# Main Header Row
ws.append(headers)
for col_idx in range(1, total_cols + 1):
    c = ws.cell(row=1, column=col_idx)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

current_row = 2

def write_section(title, test_cases):
    global current_row
    ws.cell(row=current_row, column=1, value=title)
    for col in range(1, total_cols + 1):
        c = ws.cell(row=current_row, column=col)
        c.fill = func_fill
        c.font = func_font
        c.border = border_all
    current_row += 1

    for row_data in test_cases:
        for col_idx in range(1, total_cols + 1):
            val = row_data[col_idx - 1] if col_idx - 1 < len(row_data) else ""
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = data_font
            c.border = border_all
            if col_idx in [1, 6, 9, 12]:
                c.alignment = Alignment(horizontal='left', vertical='top')
            elif col_idx in [7, 8, 10, 11, 13, 14]:
                c.alignment = Alignment(horizontal='center', vertical='top')
            else:
                c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        current_row += 1

write_section("Function 6: View Profile", test_cases_func6)
write_section("Function 7: Update Profile", test_cases_func7)
write_section("Function 12: Review Event with AI Assistance", test_cases_func12)

widths = {
    1: 18,  # TC ID
    2: 32,  # Description
    3: 36,  # Steps
    4: 38,  # Expected Result
    5: 25,  # Pre-condition
    6: 12,  # Execution 1
    7: 12,  # Tester 1
    8: 12,  # Test Date 1
    9: 12,  # Execution 2
    10: 12, # Tester 2
    11: 12, # Test Date 2
    12: 12, # Execution 3
    13: 12, # Tester 3
    14: 12  # Test Date 3
}

for col_idx, w in widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = w

output_file = r'e:\Kì 8\Event Managerment\Newest clone\EventHub\TestCases_Profile_Management.xlsx'
wb.save(output_file)
print("Excel updated successfully with comprehensive Function 12 test cases at", output_file)
